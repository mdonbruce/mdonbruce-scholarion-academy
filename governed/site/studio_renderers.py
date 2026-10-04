"""Native artifact renderers; no external requests or executable model HTML."""
import csv,io,json,textwrap

FONT_CANDIDATES=['C:/Windows/Fonts/arial.ttf','/Library/Fonts/Arial.ttf','/System/Library/Fonts/Supplemental/Arial.ttf',
 '/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf','/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',
 '/usr/share/fonts/dejavu/DejaVuSans.ttf']

def _font(size):
 """Arial where available, then common Linux/macOS fonts, then Pillow's bundled scalable font."""
 from PIL import ImageFont
 for path in FONT_CANDIDATES:
  try:return ImageFont.truetype(path,size)
  except OSError:continue
 try:return ImageFont.load_default(size=size)
 except TypeError:return ImageFont.load_default()

def render(package,outputs,edition):
 files={};missing=[]
 try:
  from PIL import Image,ImageDraw,ImageFont
  im=Image.new('RGB',(1600,900),'#063678');d=ImageDraw.Draw(im)
  font=_font(54)
  small=_font(30)
  d.text((80,80),'SCHOLARION ACADEMY',font=small,fill='#f1ca69')
  d.multiline_text((80,270),'\n'.join(textwrap.wrap(package['title'],38)),font=font,fill='white',spacing=22)
  d.text((80,720),f'Module {package["module"]} · {package["course"]}',font=small,fill='white')
  d.text((80,810),'AI-assisted content.',font=small,fill='#f1ca69')
  b=io.BytesIO();im.save(b,format='PNG');files['cover.png']=b.getvalue()
 except (ImportError,OSError):missing.append('PNG cover renderer requires Pillow')
 for index,o in enumerate(outputs):
  c=o['content'];prefix=f'{index+1:02d}-{o["kind"]}'
  if o['kind'] in ['notes','study_guide','overview']:
   try:
    from docx import Document
    doc=Document();doc.add_heading('Scholarion Academy',0);doc.add_heading(package['title'],1)
    doc.add_paragraph('AI-assisted content.')
    for line in c[edition].split('\n'):doc.add_paragraph(line)
    b=io.BytesIO();doc.save(b);files[prefix+'.docx']=b.getvalue()
   except ImportError:missing.append('DOCX renderer requires python-docx')
  if o['kind']=='slides' and isinstance(c.get('slides'),list):
   try:
    from pptx import Presentation
    from pptx.util import Inches,Pt
    deck=Presentation();deck.slide_width=Inches(13.333);deck.slide_height=Inches(7.5)
    for slide in c['slides']:
     item=deck.slides.add_slide(deck.slide_layouts[1]);item.shapes.title.text=str(slide.get('title',''))
     item.placeholders[1].text=str(slide.get('body',''))+'\n\nAI-assisted content.'
     if edition=='instructor':item.notes_slide.notes_text_frame.text=str(slide.get('notes',''))
    b=io.BytesIO();deck.save(b);files[prefix+'.pptx']=b.getvalue()
   except ImportError:missing.append('PPTX renderer requires python-pptx')
  if o['kind']=='flashcards' and isinstance(c.get('cards'),list):
   stream=io.StringIO(newline='');writer=csv.writer(stream);writer.writerow(['front','back','sourceIds'])
   cards=[]
   for card in c['cards']:
    safe={k:card.get(k,'') for k in ['front','back','sourceIds']};cards.append(safe)
    # Spreadsheet-formula injection protection for user/model text.
    writer.writerow([("'"+str(safe[k])) if str(safe[k]).startswith(('=','+','-','@')) else str(safe[k]) for k in safe])
   files[prefix+'.csv']=stream.getvalue();files[prefix+'.json']=json.dumps(cards,indent=2)
  if o['kind']=='practice_quizzes' and isinstance(c.get('questions'),list):
   questions=c['questions'] if edition=='instructor' else [{k:q.get(k) for k in ['id','prompt','options','outcomes','sourceIds']} for q in c['questions']]
   files[prefix+'.json']=json.dumps(questions,indent=2)
 files['render-status.json']=json.dumps({'edition':edition,'rendered':list(files),'unavailable':sorted(set(missing)),'narratedMedia':'not rendered; scripts require approved speech/video service'},indent=2)
 return files
