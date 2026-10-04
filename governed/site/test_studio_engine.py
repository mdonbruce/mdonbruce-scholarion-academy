import base64,io,json,tempfile,time,unittest,zipfile
from unittest.mock import patch
from governed_service import Service,Denied,Conflict
import studio_sources as sources
import studio_engine as engine
import studio_lab_runner as labs
import studio_renderers

class StudioBuildTests(unittest.TestCase):
 def setUp(self):
  self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup)
  self.s=Service(self.tmp.name);self.tokens={v['role']:k for k,v in self.s.identities.items()}
  self.p=self.source('create',title='Grounded module',module=1)
  self.p=self.source('source',id=self.p['id'],revision=1,title='Reference',text='A baseline permits comparison.',citation='Owned reference',rights='owned')
  self.sid=self.p['package']['sources'][0]['id']
  self.p=self.source('mapping',id=self.p['id'],revision=2,outcomes=[{'id':'LO1','text':'Compare a baseline','competency':'C1'}],topics=[{'id':'T1','title':'Baseline','sourceIds':[self.sid],'outcomes':['LO1']}])
 def source(self,a,**d):return sources.command(self.s,self.tokens['instructor'],'scholarion','AIM310',a,d,Denied,Conflict)
 def call(self,a,token=None,**d):return engine.command(self.s,token or self.tokens['instructor'],'scholarion','AIM310',a,d,Denied,Conflict)
 def content(self):return {'student':'Compare against a baseline ['+self.sid+']','instructor':'PRIVATE_KEY: compare the before and after results.','sourceIds':[self.sid],'supplemental':[]}
 def build(self):
  with patch.object(engine,'provider_generate',return_value=self.content()):return self.call('generate',id=self.p['id'],revision=self.p['revision'])
 def test_build_real_archives_and_edition_isolation(self):
  b=self.build();self.assertEqual(b['status'],'draft')
  result=self.call('download',build=b['build'],edition='student')
  with zipfile.ZipFile(io.BytesIO(base64.b64decode(result['base64']))) as z:
   self.assertIsNone(z.testzip());self.assertGreater(len(z.namelist()),30)
   self.assertFalse(any(b'PRIVATE_KEY' in z.read(n) for n in z.namelist()))
  result=self.call('download',build=b['build'],edition='instructor')
  with zipfile.ZipFile(io.BytesIO(base64.b64decode(result['base64']))) as z:self.assertTrue(any(b'PRIVATE_KEY' in z.read(n) for n in z.namelist()))
 def test_missing_provider_records_failure(self):
  with patch.dict('os.environ',{},clear=True):b=self.call('generate',id=self.p['id'],revision=3)
  self.assertEqual(b['status'],'incomplete');self.assertEqual(b['generated'],0);self.assertTrue(b['errors'])
 def test_edit_preserved_and_markup_escaped(self):
  c=self.content();c['student']+='<script>alert(1)</script>'
  self.p=self.source('edit',id=self.p['id'],revision=3,output='notes',body=json.dumps(c))
  b=self.call('assemble',id=self.p['id'],revision=4)
  self.assertEqual(b['generated'],1)
  result=self.call('download',build=b['build'],edition='student')
  with zipfile.ZipFile(io.BytesIO(base64.b64decode(result['base64']))) as z:
   page=z.read('01-notes.html').decode();self.assertNotIn('<script>',page);self.assertIn('&lt;script&gt;',page)
 def test_review_identity_and_version_binding(self):
  b=self.build();d={'build':b['build'],'capacity':'subject-matter-expert','decision':'approved','evidence':'Inspected all edition content and citations.'}
  with self.assertRaises(Denied):self.call('review',**d)
  for user,capacity in [('reviewer1','subject-matter-expert'),('reviewer2','instructional-designer')]:
   self.s.identities[user]={'user':user,'role':'instructor','tenant':'scholarion','expires':time.time()+100,'studioReviewRoles':[capacity]}
   with self.s.connect('scholarion') as db:db.execute('INSERT INTO enrollments VALUES(?,?,?,?)',(user,'AIM310','instructor','active'))
   self.assertTrue(self.call('review',token=user,**dict(d,capacity=capacity))['recorded'])
  self.source('edit',id=self.p['id'],revision=3,output='notes',body=json.dumps(self.content()))
  with self.assertRaises(Conflict):self.call('review',token='reviewer1',**d)
 def test_permissions(self):
  with self.assertRaises(Denied):self.call('list',token=self.tokens['student'])
  with self.assertRaises(Conflict):self.call('generate',id=self.p['id'],revision=1)
 def test_release_does_not_require_instructor_review(self):
  b=self.build();record=self.call('get',build=b['build'])
  self.assertEqual(record['reviews'],[])
  self.assertTrue(record['releasePolicy']['releaseAllowed'])
  self.assertFalse(record['releasePolicy']['instructorReviewRequired'])
  self.assertFalse(engine.release_policy({'status':'incomplete','revision':3},3)['releaseAllowed'])
  self.assertFalse(engine.release_policy({'status':'draft','revision':2},3)['releaseAllowed'])
 def test_lab_never_falls_back_to_host(self):
  with patch.dict('os.environ',{},clear=True),patch.object(labs.subprocess,'run') as run:
   self.assertEqual(labs.execute({})['status'],'blocked');run.assert_not_called()
 def test_unknown_citations_rejected(self):
  c=self.content();c['sourceIds']=['invented']
  with self.assertRaises(ValueError):engine.validate_content(c,self.p['package'])
 def test_native_renderers(self):
  try:
   from pptx import Presentation
   from docx import Document
   from PIL import Image
  except ImportError:self.skipTest('Native renderers require bundled document runtime')
  c=self.content();c['slides']=[{'title':'Baseline','body':'Compare before and after','notes':'INSTRUCTOR_NOTES_ONLY'} for _ in range(10)]
  outputs=[{'kind':'slides','content':c},{'kind':'notes','content':c}]
  files=studio_renderers.render(self.p['package'],outputs,'student')
  self.assertEqual(Image.open(io.BytesIO(files['cover.png'])).size,(1600,900))
  deck=Presentation(io.BytesIO(files['01-slides.pptx']));self.assertEqual(len(deck.slides),10)
  self.assertTrue(all('INSTRUCTOR_NOTES_ONLY' not in s.notes_slide.notes_text_frame.text for s in deck.slides))
  doc=Document(io.BytesIO(files['02-notes.docx']));self.assertFalse(any('PRIVATE_KEY' in p.text for p in doc.paragraphs))

if __name__=='__main__':unittest.main()
